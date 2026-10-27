"""Desktop Automation Tool for Aton.

Provides window focusing, text typing (Notepad, Word, Excel, browsers),
document/spreadsheet creation and opening, and keyboard shortcut simulation.
"""

import logging
import os
import subprocess
import time
from typing import Any, Dict, List, Optional
from docx import Document
import openpyxl
import pyautogui
import pygetwindow as gw
import pyperclip

logger = logging.getLogger(__name__)


def focus_application_window(app_name: str) -> Dict[str, Any]:
    """Attempt to bring a window matching app_name to the foreground.

    Args:
        app_name: Name or title fragment of the application (e.g. 'notepad', 'word', 'excel', 'chrome').

    Returns:
        Dict with status and friendly message in Portuguese.
    """
    if not app_name:
        return {
            "status": "error",
            "message": "Nenhum nome de aplicativo foi especificado.",
        }

    app_lower = app_name.lower().strip()

    # Mapping common names to window title keywords
    keyword_map = {
        "bloco de notas": "notepad",
        "word": "word",
        "excel": "excel",
        "powerpoint": "powerpoint",
        "chrome": "chrome",
        "edge": "edge",
        "brave": "brave",
        "firefox": "firefox",
        "calculadora": "calculator",
        "terminal": "powershell",
        "vscode": "visual studio code",
    }
    search_term = keyword_map.get(app_lower, app_lower)

    try:
        all_windows = gw.getAllWindows()
        for win in all_windows:
            if not win.title:
                continue
            title_lower = win.title.lower()
            if search_term in title_lower or any(part in title_lower for part in search_term.split()):
                if win.isMinimized:
                    win.restore()
                win.activate()
                time.sleep(0.2)
                logger.info(f"Window '{win.title}' focused successfully.")
                return {
                    "status": "success",
                    "message": f"Janela '{win.title}' trazida para o primeiro plano.",
                    "window_title": win.title,
                }
    except Exception as e:
        logger.debug(f"pygetwindow focus exception: {e}")

    # Fallback: attempt to open the application if not found open
    try:
        from src.tools.windows_apps import open_application
        res = open_application(app_name)
        time.sleep(0.5)
        return res
    except Exception as exc:
        logger.error(f"Failed to focus or launch '{app_name}': {exc}")
        return {
            "status": "error",
            "message": f"Não foi possível localizar ou focar na janela de '{app_name}'.",
        }


def type_text_in_active_window(
    text: str,
    app_name: Optional[str] = None,
    use_clipboard: bool = True,
    press_enter: bool = False,
) -> Dict[str, Any]:
    """Type or paste text into the active window or target application.

    Args:
        text: Text to be written.
        app_name: Optional application name to focus first (e.g. 'notepad', 'word', 'excel', 'chrome').
        use_clipboard: If True, uses clipboard injection (preserves accents, symbols, emojis, and line breaks).
        press_enter: If True, simulates pressing the Enter key after typing.

    Returns:
        Dict with status and confirmation message.
    """
    if not text:
        return {
            "status": "error",
            "message": "Nenhum texto informado para digitação.",
        }

    # Focus application window if specified
    if app_name:
        focus_application_window(app_name)
        time.sleep(0.25)

    try:
        if use_clipboard:
            # Store original clipboard if possible
            try:
                original = pyperclip.paste()
            except Exception:
                original = ""

            pyperclip.copy(text)
            time.sleep(0.05)
            # Paste with Ctrl+V
            pyautogui.hotkey("ctrl", "v")
            time.sleep(0.05)

            # Restore original clipboard after short delay
            if original:
                pyperclip.copy(original)
        else:
            pyautogui.write(text, interval=0.01)

        if press_enter:
            time.sleep(0.05)
            pyautogui.press("enter")

        dest_str = f" no {app_name}" if app_name else " na janela ativa"
        logger.info(f"Text typed successfully{dest_str}: '{text[:50]}...'")
        return {
            "status": "success",
            "message": f"Texto inserido com sucesso{dest_str}.",
            "text_length": len(text),
        }

    except Exception as e:
        logger.error(f"Error typing text: {e}", exc_info=True)
        return {
            "status": "error",
            "message": f"Falha ao digitar o texto: {str(e)}",
        }


def create_and_open_document(
    file_name: str,
    content: str,
    file_type: str = "txt",
    directory: Optional[str] = None,
) -> Dict[str, Any]:
    """Create a structured document or spreadsheet and open it in the default Windows application.

    Args:
        file_name: Base name of the file (e.g. 'reuniao', 'financeiro').
        content: Text content or table data (separated by newlines, commas, or semicolons).
        file_type: Extension/type: 'txt', 'docx' (Word), 'xlsx' (Excel), 'csv', 'md'.
        directory: Optional directory path (defaults to user's Documents folder).

    Returns:
        Dict with status, message, file path, and format info.
    """
    file_type = (file_type or "txt").lower().replace(".", "").strip()
    clean_name = os.path.splitext(file_name)[0].strip() or "documento_aton"
    full_filename = f"{clean_name}.{file_type}"

    # Default to Documents folder
    if not directory:
        directory = os.path.expanduser("~/Documents")
    os.makedirs(directory, exist_ok=True)
    filepath = os.path.join(directory, full_filename)

    try:
        # 1. Word Document (.docx)
        if file_type == "docx":
            doc = Document()
            lines = content.strip().split("\n")
            if lines:
                doc.add_heading(lines[0], level=1)
                for line in lines[1:]:
                    clean_line = line.strip()
                    if clean_line:
                        doc.add_paragraph(clean_line)
            doc.save(filepath)

        # 2. Excel Spreadsheet (.xlsx)
        elif file_type in ("xlsx", "excel"):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Dados"
            rows = content.strip().split("\n")
            for r_idx, row in enumerate(rows, start=1):
                # Detect separator
                if ";" in row:
                    cols = [c.strip() for c in row.split(";")]
                elif "\t" in row:
                    cols = [c.strip() for c in row.split("\t")]
                elif "|" in row:
                    cols = [c.strip() for c in row.split("|") if c.strip()]
                elif "," in row:
                    cols = [c.strip() for c in row.split(",")]
                else:
                    cols = [row.strip()]

                for c_idx, val in enumerate(cols, start=1):
                    # Convert to number if numeric
                    if val.isdigit():
                        cell_val = int(val)
                    else:
                        try:
                            cell_val = float(val)
                        except ValueError:
                            cell_val = val
                    ws.cell(row=r_idx, column=c_idx, value=cell_val)
            wb.save(filepath)

        # 3. Text, Markdown, CSV (.txt, .md, .csv)
        else:
            with open(filepath, "w", encoding="utf-8") as f:
                f.write(content)

        # Automatically open file with default Windows program
        try:
            os.startfile(filepath)
            logger.info(f"Opened document with default application: {filepath}")
        except Exception as open_err:
            logger.warning(f"File created at '{filepath}' but could not auto-open: {open_err}")

        app_name_map = {
            "docx": "Microsoft Word",
            "xlsx": "Microsoft Excel",
            "csv": "Excel / Editor de Planilhas",
            "txt": "Bloco de Notas",
            "md": "Editor Markdown",
        }
        app_desc = app_name_map.get(file_type, "aplicativo padrão")

        return {
            "status": "success",
            "message": f"Documento '{full_filename}' criado com sucesso e aberto no {app_desc}.",
            "file_path": filepath,
            "file_type": file_type,
        }

    except Exception as e:
        logger.error(f"Error creating and opening document '{filepath}': {e}", exc_info=True)
        return {
            "status": "error",
            "message": f"Não foi possível criar o documento: {str(e)}",
        }


def send_keyboard_shortcut(
    shortcut: str,
    app_name: Optional[str] = None,
) -> Dict[str, Any]:
    """Execute keyboard shortcuts such as 'ctrl+s', 'ctrl+a', 'ctrl+c', 'enter', 'tab'.

    Args:
        shortcut: Key combination string (e.g. 'ctrl+s', 'ctrl+c', 'enter', 'tab', 'esc').
        app_name: Optional application name to focus before pressing keys.

    Returns:
        Dict with status and confirmation message.
    """
    if not shortcut:
        return {
            "status": "error",
            "message": "Nenhum atalho de teclado informado.",
        }

    if app_name:
        focus_application_window(app_name)
        time.sleep(0.2)

    normalized = shortcut.lower().replace(" ", "").replace("-", "+")
    keys = normalized.split("+")

    try:
        if len(keys) == 1:
            pyautogui.press(keys[0])
        else:
            pyautogui.hotkey(*keys)

        logger.info(f"Keyboard shortcut executed: {normalized}")
        return {
            "status": "success",
            "message": f"Atalho '{normalized}' executado com sucesso.",
            "shortcut": normalized,
        }
    except Exception as e:
        logger.error(f"Error executing shortcut '{shortcut}': {e}", exc_info=True)
        return {
            "status": "error",
            "message": f"Falha ao executar o atalho '{shortcut}': {str(e)}",
        }


# Declarations for Gemini Live API
DESKTOP_AUTOMATION_FUNCTION_DECLARATIONS = [
    {
        "name": "type_text_in_active_window",
        "description": "Digita ou cola texto diretamente na janela ativa ou no software indicado (Bloco de Notas, Word, Excel, navegador, formulários da web), com suporte a acentuação e quebras de linha.",
        "parameters": {
            "type": "OBJECT",
            "properties": {
                "text": {
                    "type": "STRING",
                    "description": "O texto a ser digitado ou colado no aplicativo.",
                },
                "app_name": {
                    "type": "STRING",
                    "description": "Nome opcional do aplicativo para focar antes de digitar (ex: 'notepad', 'word', 'excel', 'chrome', 'edge').",
                },
                "press_enter": {
                    "type": "BOOLEAN",
                    "description": "Se verdadeiro, pressiona a tecla Enter após digitar o texto.",
                },
            },
            "required": ["text"],
        },
    },
    {
        "name": "create_and_open_document",
        "description": "Cria um novo documento de texto, documento do Word ou planilha do Excel com o conteúdo desejado e abre instantaneamente no programa correspondente.",
        "parameters": {
            "type": "OBJECT",
            "properties": {
                "file_name": {
                    "type": "STRING",
                    "description": "Nome do arquivo (ex: 'anotacoes', 'lista_compras', 'relatorio_trimestral').",
                },
                "content": {
                    "type": "STRING",
                    "description": "Conteúdo do documento. Para planilhas, separe colunas por ponto e vírgula, vírgula ou tabulação.",
                },
                "file_type": {
                    "type": "STRING",
                    "description": "Tipo do arquivo: 'txt' (Bloco de Notas), 'docx' (Word), 'xlsx' (Excel), 'csv' ou 'md'. Padrão é 'txt'.",
                },
            },
            "required": ["file_name", "content"],
        },
    },
    {
        "name": "send_keyboard_shortcut",
        "description": "Envia um atalho de teclado para o Windows ou para o aplicativo em foco (ex: 'ctrl+s' para salvar, 'ctrl+a' para selecionar tudo, 'ctrl+c', 'ctrl+v', 'enter', 'tab', 'esc').",
        "parameters": {
            "type": "OBJECT",
            "properties": {
                "shortcut": {
                    "type": "STRING",
                    "description": "Combinação de teclas a enviar (ex: 'ctrl+s', 'ctrl+a', 'enter', 'tab').",
                },
                "app_name": {
                    "type": "STRING",
                    "description": "Nome opcional do aplicativo para focar antes do atalho (ex: 'notepad', 'word', 'excel').",
                },
            },
            "required": ["shortcut"],
        },
    },
    {
        "name": "focus_application_window",
        "description": "Traz a janela de um aplicativo específico para o primeiro plano na tela (ex: 'notepad', 'word', 'excel', 'chrome').",
        "parameters": {
            "type": "OBJECT",
            "properties": {
                "app_name": {
                    "type": "STRING",
                    "description": "Nome do programa para focar na tela.",
                },
            },
            "required": ["app_name"],
        },
    },
]


def dispatch_desktop_automation_tool(tool_name: str, tool_args: Dict[str, Any]) -> Dict[str, Any]:
    """Dispatch desktop typing, document, and keyboard automation tools."""
    if tool_name == "type_text_in_active_window":
        text = tool_args.get("text", "")
        app_name = tool_args.get("app_name")
        press_enter = bool(tool_args.get("press_enter", False))
        return type_text_in_active_window(text=text, app_name=app_name, press_enter=press_enter)

    elif tool_name == "create_and_open_document":
        file_name = tool_args.get("file_name", "documento")
        content = tool_args.get("content", "")
        file_type = tool_args.get("file_type", "txt")
        directory = tool_args.get("directory")
        return create_and_open_document(
            file_name=file_name,
            content=content,
            file_type=file_type,
            directory=directory,
        )

    elif tool_name == "send_keyboard_shortcut":
        shortcut = tool_args.get("shortcut", "")
        app_name = tool_args.get("app_name")
        return send_keyboard_shortcut(shortcut=shortcut, app_name=app_name)

    elif tool_name == "focus_application_window":
        app_name = tool_args.get("app_name", "")
        return focus_application_window(app_name=app_name)

    return {
        "status": "error",
        "message": f"Ferramenta de automação de desktop desconhecida '{tool_name}'.",
    }
